"""CSV/TSV/XLSX comment records with metadata and explicit text columns."""
from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

TEXT_FIELD = "text"
LABEL_FIELD = "label"
CSV_SUFFIXES = (".csv", ".tsv")
EXCEL_SUFFIXES = (".xlsx",)
TEXT_COLUMNS = ("text", "text_raw", "comment", "content", "评论", "评论内容", "内容")
LABEL_ALIASES = {
    "0": "negative", "1": "positive", "正面": "positive", "负面": "negative",
    "中性": "neutral", "不确定": "uncertain", "无关": "irrelevant",
}
METADATA_ALIASES = {
    "comment_id": ("rpid",),
    "video_id": ("bvid",),
    "posted_at": ("ctime_utc",),
    "collected_at": ("collected_at_utc",),
}
logger = logging.getLogger(__name__)


@dataclass
class Record:
    text: str
    label: str | None = None
    tokens: str | None = None
    meta: dict[str, Any] = field(default_factory=dict)


def detect_text_column(columns: list[str]) -> str:
    """Never silently select an ID in a multi-column file."""
    lookup = {str(name).strip().casefold(): str(name) for name in columns}
    for candidate in TEXT_COLUMNS:
        if candidate in lookup:
            return lookup[candidate]
    if len(columns) == 1 and next(iter(lookup)) not in {"id", "comment_id", "video_id", "label"}:
        return str(columns[0])
    raise ValueError("无法识别评论文本列，请指定 --text-column 或使用 text/text_raw/评论 列")


def _read_rows(path: Path, sheet_name: str | None) -> list[list[Any]]:
    if path.suffix.lower() in CSV_SUFFIXES:
        content = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("文件编码无法识别，请转换为 UTF-8")
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        return list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        raise ValueError("仅支持 .csv、.tsv、.xlsx 文件")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("读取 XLSX 需要 openpyxl：python -m pip install openpyxl") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, BadZipFile) as exc:
        raise ValueError(f"无法读取 XLSX 文件：{path.name}") from exc
    try:
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            raise ValueError(f"工作表不存在：{sheet_name}")
        sheet = workbook[sheet_name] if sheet_name is not None else workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def load_records(
    path: str | Path, *, text_column: str | None = None,
    sheet_name: str | None = None, keep_empty: bool = False,
) -> list[Record]:
    """Preserve auxiliary columns. input_row counts the header as row 1.

    keep_empty=True lets the preparation pipeline audit invalid records.
    """
    path = Path(path)
    rows = _read_rows(path, sheet_name)
    if not rows:
        raise ValueError("数据文件为空")
    header = [str(value).strip() if value is not None else "" for value in rows[0]]
    if not all(header) or len({value.casefold() for value in header}) != len(header):
        raise ValueError("表头不能包含空列名或重复列名")
    lookup = {name.casefold(): name for name in header}
    if text_column is not None:
        if text_column.strip().casefold() not in lookup:
            raise ValueError(f"文本列不存在：{text_column}")
        selected = lookup[text_column.strip().casefold()]
    else:
        selected = detect_text_column(header)
    label_key, tokens_key = lookup.get("label"), lookup.get("tokens")
    records, skipped = [], 0
    for row_number, row in enumerate(rows[1:], 2):
        if len(row) > len(header) and any(value not in (None, "") for value in row[len(header):]):
            raise ValueError(f"第 {row_number} 行列数超过表头，请检查 CSV 引号或分隔符")
        values = dict(zip(header, row + [None] * max(0, len(header) - len(row))))
        raw_text = values[selected]
        text = raw_text if isinstance(raw_text, str) else ""
        if not text.strip() and not keep_empty:
            skipped += 1
            continue
        raw_label = values.get(label_key) if label_key else None
        label = str(raw_label).strip().lower() if raw_label is not None else ""
        label = LABEL_ALIASES.get(label, label) or None
        standard_meta = {"id", "comment_id", "video_id", "source", "source_ref", "collected_at",
                         "crawled_at", "posted_at", "input_row", "input_file", "rpid", "bvid",
                         "ctime_utc", "collected_at_utc"}
        meta = {key.casefold() if key.casefold() in standard_meta else key: value
                for key, value in values.items() if key not in {selected, label_key, tokens_key}}
        for target, aliases in METADATA_ALIASES.items():
            if meta.get(target) in (None, ""):
                for alias in aliases:
                    if meta.get(alias) not in (None, ""):
                        meta[target] = meta[alias]
                        break
        for key in ("id", "comment_id", "video_id", "rpid", "bvid"):
            if key in meta and meta[key] is not None:
                meta[key] = str(meta[key]).strip()
        meta.setdefault("input_row", row_number)
        meta.setdefault("input_file", path.name)
        if raw_text is not None and not isinstance(raw_text, str):
            meta["invalid_text_type"] = type(raw_text).__name__
        tokens = values.get(tokens_key) if tokens_key else None
        records.append(Record(text, label, str(tokens) if tokens is not None else None, meta))
    if skipped:
        logger.warning("%s: 跳过 %d 条空白或非文本评论", path.name, skipped)
    return records


def save_records(records: list[Record], path: str | Path) -> Path:
    """Write BOM CSV, retaining metadata needed for lineage and grouping."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [TEXT_FIELD, LABEL_FIELD]
    if any(record.tokens is not None for record in records):
        fields.append("tokens")
    fields += sorted({key for record in records for key in record.meta} - {"text", "label", "tokens"})
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for record in records:
            row = {key: value for key, value in record.meta.items() if key in fields}
            row.update(text=record.text, label=record.label or "")
            if "tokens" in fields:
                row["tokens"] = record.tokens or ""
            writer.writerow(row)
    return path
